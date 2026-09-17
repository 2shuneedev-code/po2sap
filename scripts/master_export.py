#!/usr/bin/env python
"""masters/ → 엑셀 한 파일. 현업에게 줄 편집 양식을 만든다.

    python scripts/master_export.py                      # masters/거래처마스터.xlsx
    python scripts/master_export.py --out /경로/파일.xlsx

**항상 여기서 뽑은 최신 파일로 편집한다.** 예전 엑셀을 고쳐서 가져오면 그 사이
다른 사람이 바꾼 것을 덮어쓴다. 가져오기가 그런 경우를 잡아내지 못한다 —
엑셀에는 "언제 뽑았는지"만 적혀 있고 무엇이 바뀌었는지는 모른다.

LLM 호출 없음 = 비용 0.
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

import yaml  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

sys.path.insert(0, str(ROOT / "scripts"))
from master_sheets import (  # noqa: E402
    ALLOWED,
    COLUMNS,
    GUIDE,
    READ_ONLY,
    S_BRAND,
    S_CUSTOMER,
    S_DOC,
    S_EXPR,
    S_FIXED,
    S_GUIDE,
    S_HINTS,
    S_TABLE,
    file_types_to_text,
)

HEADER_FILL = PatternFill("solid", fgColor="1F5FD0")
LOCKED_FILL = PatternFill("solid", fgColor="EDEFF2")
TODO_FILL = PatternFill("solid", fgColor="FFF5E0")


# ── masters/ 읽기 ─────────────────────────────────────────────────────
MASTERS = ROOT / "masters"        # --masters 로 바꾼다 (테스트가 실물을 안 건드리게)


def customer_files() -> list[Path]:
    return [p for p in sorted((MASTERS / "customers").glob("*.yaml"))
            if not p.stem.startswith("_")]


def load(path: Path) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8")) or {}


def read_csv(rel: str) -> list[dict]:
    import csv
    path = MASTERS / rel
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def collect() -> dict[str, list[list]]:
    """시트별 행 목록을 만든다."""
    rows: dict[str, list[list]] = {name: [] for name in COLUMNS}
    brand_names = {
        (r.get("kunnr", ""), r.get("zbrand", "")): r.get("zbrant", "")
        for r in read_csv("refs/brand_master.csv")
    }

    for path in customer_files():
        d = load(path)
        meta = d.get("meta") or {}
        code = str(meta.get("code") or path.stem.upper())
        split = d.get("split") or {}

        rows[S_CUSTOMER].append([
            code, meta.get("name", ""), str(meta.get("customer_no", "")),
            file_types_to_text(meta.get("file_types") or []),
            meta.get("owner", ""), meta.get("status", "draft"),
            split.get("by", "none"), split.get("label", ""),
        ])

        for field, spec in (d.get("fields") or {}).items():
            if not isinstance(spec, dict):
                continue
            source = spec.get("from")
            if source == "const":
                rows[S_FIXED].append([
                    code, field, str(spec.get("value", "")),
                    spec.get("todo", ""), spec.get("explain", ""),
                ])
            elif source == "doc":
                rows[S_DOC].append([
                    code, field, spec.get("path", ""), spec.get("format", ""),
                    "Y" if spec.get("required") else "N", spec.get("fallback", ""),
                ])
            elif source == "expr":
                rows[S_EXPR].append([
                    code, field, spec.get("expr", ""), spec.get("explain", ""),
                ])

        for table_id, table in (d.get("tables") or {}).items():
            if not isinstance(table, dict):
                continue
            conds = [c for c in (table.get("when") or []) if isinstance(c, dict)]
            results = list(table.get("then") or [])
            no_match = table.get("on_no_match") or {}
            for row in table.get("rows") or []:
                if not isinstance(row, dict):
                    continue
                whens = list(row.get("when") or [])
                thens = list(row.get("then") or [])
                rows[S_TABLE].append([
                    code, table_id, table.get("scope", "header"),
                    " / ".join(str(c.get("source", "")) for c in conds),
                    " / ".join(str(c.get("op", "")) for c in conds),
                    " / ".join(str(w) for w in whens),
                    " / ".join(str(r) for r in results),
                    " / ".join(str(t) for t in thens),
                    f"{no_match.get('action', '')}: {no_match.get('message', '')}".strip(": "),
                ])

        hints = (d.get("extraction") or {}).get("hints") or ""
        if hints:
            rows[S_HINTS].append([code, hints])

    for key in read_csv("refs/brand_keys.csv"):
        kunnr, zbrand = key.get("kunnr", ""), key.get("zbrand", "")
        rows[S_BRAND].append([
            kunnr, zbrand, brand_names.get((kunnr, zbrand), ""),
            key.get("text", ""), key.get("match", "contains"), key.get("note", ""),
        ])

    return rows


# ── 엑셀 쓰기 ─────────────────────────────────────────────────────────
def write(rows: dict[str, list[list]], out: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    guide = wb.create_sheet(S_GUIDE)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 96
    guide["A1"] = "거래처 마스터"
    guide["A1"].font = Font(bold=True, size=14)
    guide["B2"] = f"뽑은 시각: {datetime.now():%Y-%m-%d %H:%M}"
    guide["B2"].font = Font(color="808080")
    for index, (left, right) in enumerate(GUIDE, start=4):
        guide.cell(index, 1, left).font = Font(bold=bool(left) and not right)
        cell = guide.cell(index, 2, right)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for name, columns in COLUMNS.items():
        ws = wb.create_sheet(name)
        read_only = name in READ_ONLY

        for col, title in enumerate(columns, start=1):
            cell = ws.cell(1, col, title)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")

        for r, row in enumerate(rows.get(name, []), start=2):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                cell.alignment = Alignment(vertical="top", wrap_text=name == S_HINTS)
                if read_only:
                    cell.fill = LOCKED_FILL
                elif columns[c - 1] == "확인필요(todo)" and value:
                    cell.fill = TODO_FILL

        _size(ws, name, columns, rows.get(name, []))
        _dropdowns(ws, name, columns, len(rows.get(name, [])))
        ws.freeze_panes = "A2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def _size(ws, name: str, columns: list[str], rows: list[list]) -> None:
    for index, title in enumerate(columns, start=1):
        letter = get_column_letter(index)
        if name == S_HINTS and title == "읽기 힌트":
            ws.column_dimensions[letter].width = 110
            continue
        widest = max([len(str(title))] + [len(str(r[index - 1])) for r in rows if len(r) >= index],
                     default=10)
        ws.column_dimensions[letter].width = min(max(widest + 3, 10), 46)
    if name == S_HINTS:
        for r in range(2, len(rows) + 2):
            ws.row_dimensions[r].height = 200


def _dropdowns(ws, name: str, columns: list[str], count: int) -> None:
    """허용값이 정해진 칸은 드롭다운으로 — 오타가 가져오기에서 막히기 전에 막는다."""
    for index, title in enumerate(columns, start=1):
        options = ALLOWED.get((name, title))
        if not options:
            continue
        dv = DataValidation(
            type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
            showErrorMessage=True, errorTitle="허용되지 않는 값",
            error="목록에서 고르세요: " + ", ".join(options),
        )
        ws.add_data_validation(dv)
        letter = get_column_letter(index)
        dv.add(f"{letter}2:{letter}{max(count + 1, 200)}")


def main() -> int:
    parser = argparse.ArgumentParser(description="masters/ → 거래처 마스터 엑셀")
    parser.add_argument("--out", default="")
    parser.add_argument("--masters", default="",
                        help="마스터 폴더 (기본 masters/). 테스트가 사본을 가리킬 때 쓴다")
    args = parser.parse_args()

    global MASTERS
    if args.masters:
        MASTERS = Path(args.masters)
    out = Path(args.out) if args.out else MASTERS / "거래처마스터.xlsx"

    rows = collect()
    write(rows, out)

    print(f"만들었습니다: {out}")
    for name in COLUMNS:
        mark = " (읽기전용)" if name in READ_ONLY else ""
        print(f"  {name:20} {len(rows.get(name, [])):4}행{mark}")
    print("\n고친 뒤:  python scripts/master_import.py " + str(out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
