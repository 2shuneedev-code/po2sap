"""SALES ORDER 템플릿이 전송 필드 목록·순서를 정한다 (`masters/template.py`).

템플릿 2행 = 필드 코드, 1행 = 영문 머리글. 열을 넣고 빼고 옮기면 검수 표·전송이 따라온다.
"""

from __future__ import annotations

import shutil

import openpyxl
import pytest
from app.masters.loader import load_customer
from app.masters.template import template_columns

TEMPLATE = "templates/SALES ORDER.xlsx"


@pytest.fixture
def workspace(tmp_path, masters_dir):
    dst = tmp_path / "masters"
    shutil.copytree(masters_dir, dst)
    return dst


def _edit(workspace, fn) -> None:
    path = workspace / TEMPLATE
    wb = openpyxl.load_workbook(path)
    fn(wb.active)
    wb.save(path)


def test_field_order_is_the_template_order(workspace):
    codes = [c for c, _ in template_columns(workspace / TEMPLATE)]
    assert list(load_customer("kl", workspace).raw["field_specs"]) == codes


def test_headers_come_from_row_one(workspace):
    specs = load_customer("kl", workspace).raw["field_specs"]
    for code, head in template_columns(workspace / TEMPLATE):
        assert specs[code]["sheet"] == head


def test_new_template_column_is_sent_blank_with_a_todo(workspace):
    def add(ws):
        col = ws.max_column + 1
        ws.cell(1, col, "NEW THING")
        ws.cell(2, col, "ZNEW")

    _edit(workspace, add)
    master = load_customer("kl", workspace)
    assert list(master.raw["field_specs"])[-1] == "ZNEW"
    assert master.raw["field_specs"]["ZNEW"]["label"] == "NEW THING"
    assert master.fields["ZNEW"]["from"] == "const" and master.fields["ZNEW"]["value"] == ""
    assert "todo" in master.fields["ZNEW"]


def test_removed_template_column_is_not_sent(workspace):
    def drop_price(ws):
        for col in range(ws.max_column, 0, -1):
            if str(ws.cell(2, col).value or "").strip() == "PRICE":
                ws.delete_cols(col)

    _edit(workspace, drop_price)
    master = load_customer("kl", workspace)
    assert "PRICE" not in master.raw["field_specs"]
    assert "PRICE" not in master.fields


def test_real_template_matches_the_attribute_store():
    """실물 템플릿의 필드는 전부 sap_defaults 에 속성(라벨 등)이 있다 — 새로 넣었으면 속성도 채운다."""
    from pathlib import Path

    import yaml

    root = Path(__file__).resolve().parents[2] / "masters"
    specs = yaml.safe_load((root / "_base" / "sap_defaults.yaml").read_text("utf-8"))["field_specs"]
    missing = [c for c, _ in template_columns(root / TEMPLATE) if c not in specs]
    assert not missing, f"sap_defaults.yaml field_specs 에 속성이 없습니다: {missing}"
