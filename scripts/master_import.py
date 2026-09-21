#!/usr/bin/env python
"""엑셀 → masters/. 현업이 고친 값을 검증하고 반영한다.

    python scripts/master_import.py masters/거래처마스터.xlsx           # 계획 보고 확인
    python scripts/master_import.py masters/거래처마스터.xlsx --dry-run  # 계획만
    python scripts/master_import.py masters/거래처마스터.xlsx --yes      # 확인 없이

**먼저 전부 검사하고, 하나라도 걸리면 아무것도 쓰지 않는다.** 절반만 반영되면
어느 거래처가 어느 상태인지 아무도 모르게 된다.

YAML 은 ruamel 로 **주석·읽기힌트를 보존한 채** 값만 갈아끼운다. 산출식(expr)과
힌트는 엑셀에서 읽기전용이라 손대지 않는다.

LLM 호출 없음 = 비용 0.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT / "scripts"))

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _console import use_utf8  # noqa: E402

use_utf8()   # 윈도우(cp949)에서 파이프로 넘길 때 한글·— 가 죽지 않게

from app.masters import brands as brand_store  # noqa: E402
from master_sheets import (  # noqa: E402
    ALLOWED,
    COLUMNS,
    READ_ONLY,
    S_BRAND,
    S_CUSTOMER,
    S_DOC,
    S_FIXED,
    as_text,
    text_to_file_types,
)
from openpyxl import load_workbook  # noqa: E402
from ruamel.yaml import YAML  # noqa: E402

MASTERS = ROOT / "masters"        # --masters 로 바꾼다 (테스트가 실물을 안 건드리게)
CUSTOMERS = MASTERS / "customers"


def yaml_rt() -> YAML:
    y = YAML(typ="rt")
    y.preserve_quotes = True
    y.width = 4096                 # 긴 hints 가 멋대로 접히지 않게
    y.indent(mapping=2, sequence=4, offset=2)
    return y


# ── 엑셀 읽기 ─────────────────────────────────────────────────────────
def read_sheets(path: Path) -> dict[str, list[dict]]:
    wb = load_workbook(path, data_only=True)
    out: dict[str, list[dict]] = {}
    for name, columns in COLUMNS.items():
        if name in READ_ONLY or name not in wb.sheetnames:
            continue
        ws = wb[name]
        header = [as_text(c.value) for c in ws[1]]
        rows = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            record = {header[i]: as_text(v) for i, v in enumerate(raw) if i < len(header)}
            if any(record.get(c) for c in columns):      # 빈 줄 건너뛰기
                rows.append(record)
        out[name] = rows
    return out


# ── 검증 ─────────────────────────────────────────────────────────────
def field_names() -> list[str]:
    y = yaml_rt()
    base = y.load((MASTERS / "_base" / "sap_defaults.yaml").read_text(encoding="utf-8"))
    return list(base.get("field_specs") or {})


def validate(sheets: dict[str, list[dict]]) -> list[str]:
    errors: list[str] = []
    fields = set(field_names())
    codes = {r["거래처코드"] for r in sheets.get(S_CUSTOMER, []) if r.get("거래처코드")}

    if not codes:
        errors.append(f"[{S_CUSTOMER}] 거래처가 한 곳도 없습니다.")

    seen: set[str] = set()
    for index, row in enumerate(sheets.get(S_CUSTOMER, []), start=2):
        code = row.get("거래처코드", "")
        where = f"[{S_CUSTOMER}] {index}행 {code or '(코드없음)'}"
        if not code:
            errors.append(f"{where} — 거래처코드가 비었습니다.")
        elif code in seen:
            errors.append(f"{where} — 거래처코드가 중복입니다.")
        seen.add(code)
        if not row.get("고객코드"):
            errors.append(f"{where} — 고객코드가 비었습니다. 참조표 조회가 전부 빗나갑니다.")
        for column in ("상태", "오더분할"):
            allowed = ALLOWED.get((S_CUSTOMER, column), [])
            value = row.get(column, "")
            if value and value not in allowed:
                errors.append(f"{where} — {column} '{value}' 는 허용되지 않습니다 ({', '.join(allowed)}).")

    for sheet, column in ((S_FIXED, "전송필드"), (S_DOC, "전송필드")):
        for index, row in enumerate(sheets.get(sheet, []), start=2):
            code, field = row.get("거래처코드", ""), row.get(column, "")
            where = f"[{sheet}] {index}행"
            if code not in codes:
                errors.append(f"{where} — '{code}' 는 {S_CUSTOMER} 시트에 없는 거래처입니다.")
            if field and field not in fields:
                errors.append(
                    f"{where} — '{field}' 는 전송 필드가 아닙니다 "
                    "(masters/_base/sap_defaults.yaml 에 있는 이름만 됩니다)."
                )

    errors += _validate_brands(sheets.get(S_BRAND, []))
    return errors


def _validate_brands(rows: list[dict]) -> list[str]:
    """그 고객에 SAP 이 등록한 코드인지 · 같은 문구가 두 코드에 가지 않는지."""
    errors: list[str] = []
    grouped: dict[tuple[str, str], list] = {}

    for index, row in enumerate(rows, start=2):
        kunnr, zbrand = row.get("고객코드", ""), row.get("브랜드코드", "")
        text, match = row.get("발주서 원문 문구", ""), row.get("비교", "contains")
        if not (kunnr and zbrand and text):
            if kunnr or zbrand or text:
                errors.append(f"[{S_BRAND}] {index}행 — 고객코드·브랜드코드·원문 문구는 모두 필요합니다.")
            continue
        if match not in ALLOWED[(S_BRAND, "비교")]:
            errors.append(f"[{S_BRAND}] {index}행 — 비교 '{match}' 는 허용되지 않습니다.")
            continue
        grouped.setdefault((kunnr, zbrand), []).append(
            brand_store.BrandKey(kunnr=kunnr, zbrand=zbrand, match=match,
                                 text=text, note=row.get("비고", ""))
        )

    for (kunnr, zbrand), keys in grouped.items():
        try:
            # 다른 코드가 쓰는 문구인지도 여기서 걸린다 — 저장 로직과 같은 검사다.
            brand_store.validate_keys(MASTERS, kunnr, zbrand, keys)
        except (brand_store.BrandError, ValueError) as exc:
            errors.append(f"[{S_BRAND}] 고객 {kunnr} / 코드 {zbrand} — {exc}")
    return errors


# ── 계획 ─────────────────────────────────────────────────────────────
def plan(sheets: dict[str, list[dict]]) -> tuple[dict[str, dict], list[str]]:
    """거래처별로 바꿀 내용과, 사람이 확인해야 할 경고를 만든다."""
    warnings: list[str] = []
    changes: dict[str, dict] = {}

    for row in sheets.get(S_CUSTOMER, []):
        code = row["거래처코드"]
        changes[code] = {
            "meta": {
                "name": row.get("거래처명", ""),
                "customer_no": row.get("고객코드", ""),
                "owner": row.get("담당자", ""),
                "status": row.get("상태", "draft") or "draft",
                "file_types": text_to_file_types(row.get("파일형식", "")),
            },
            "split": {"by": row.get("오더분할", "none") or "none",
                      "label": row.get("분할설명", "")},
            "fixed": {}, "doc": {},
        }

    for row in sheets.get(S_FIXED, []):
        code, field = row.get("거래처코드"), row.get("전송필드")
        if code in changes and field:
            changes[code]["fixed"][field] = {
                "value": row.get("값", ""),
                "todo": row.get("확인필요(todo)", ""),
                "explain": row.get("설명", ""),
            }

    for row in sheets.get(S_DOC, []):
        code, field = row.get("거래처코드"), row.get("전송필드")
        if code in changes and field:
            changes[code]["doc"][field] = {
                "path": row.get("문서경로", ""),
                "format": row.get("형식", ""),
                "required": row.get("필수", "").upper() == "Y",
                "fallback": row.get("대체경로", ""),
            }

    # 엑셀에서 사라진 필드는 **덮어쓰기를 그만둔다**는 뜻이다. 되돌리기 어려우니 알린다.
    y = yaml_rt()
    for code, change in changes.items():
        path = CUSTOMERS / f"{code.lower()}.yaml"
        if not path.exists():
            warnings.append(f"{code} — 새 거래처입니다. `{path.name}` 을 템플릿에서 만듭니다.")
            continue
        current = y.load(path.read_text(encoding="utf-8")) or {}
        kept = set(change["fixed"]) | set(change["doc"])
        for field, spec in (current.get("fields") or {}).items():
            if not isinstance(spec, dict):
                continue
            if spec.get("from") in ("const", "doc") and field not in kept:
                warnings.append(
                    f"{code}.{field} — 엑셀에서 빠졌습니다. 이 거래처의 덮어쓰기를 지우고 "
                    "표준 프로필 값으로 되돌립니다."
                )
    return changes, warnings


# ── 반영 ─────────────────────────────────────────────────────────────
def _flow(mapping: dict):
    """`{from: const, value: "A"}` 처럼 한 줄로 쓰는 형태를 유지한다."""
    from ruamel.yaml.comments import CommentedMap

    out = CommentedMap(mapping)
    out.fa.set_flow_style()
    return out


def _set(container, key, value) -> bool:
    """값이 **실제로 다를 때만** 쓴다. 바뀐 게 있으면 True.

    무조건 대입하면 ruamel 이 그 자리를 다시 찍어내면서 들여쓰기·flow 스타일이
    바뀐다. 아무것도 안 고친 가져오기가 70줄짜리 diff 를 내면 진짜 규칙 변경이
    그 안에 묻힌다 — Git 이력이 곧 규칙 변경 대장인데(CLAUDE.md §6) 못 읽게 된다.
    """
    if key in container and container[key] == value:
        return False
    container[key] = value
    return True


def _field_entry(spec: dict, kind: str) -> dict:
    if kind == "const":
        entry = {"from": "const", "value": spec["value"]}
        if spec["todo"]:
            entry["todo"] = spec["todo"]
        if spec["explain"]:
            entry["explain"] = spec["explain"]
        return entry
    entry = {"from": "doc", "path": spec["path"]}
    if spec["format"]:
        entry["format"] = spec["format"]
    if spec["required"]:
        entry["required"] = True
    if spec["fallback"]:
        entry["fallback"] = spec["fallback"]
    return entry


def apply(changes: dict[str, dict], sheets: dict[str, list[dict]],
          *, normalize: bool = False) -> list[str]:
    y = yaml_rt()
    touched: list[str] = []

    for code, change in changes.items():
        path = CUSTOMERS / f"{code.lower()}.yaml"
        fresh = not path.exists()
        source = (CUSTOMERS / "_template.yaml") if fresh else path
        doc = y.load(source.read_text(encoding="utf-8"))
        dirty = fresh or normalize

        meta = doc.setdefault("meta", {})
        dirty |= _set(meta, "code", code)
        for key, value in change["meta"].items():
            if value in (None, "") and key not in ("name", "owner"):
                continue
            dirty |= _set(meta, key, value)

        split = doc.setdefault("split", {})
        dirty |= _set(split, "by", change["split"]["by"])
        if change["split"]["label"]:
            dirty |= _set(split, "label", change["split"]["label"])

        fields = doc.setdefault("fields", {})
        # 엑셀에 없는 const/doc 필드는 덮어쓰기를 거둔다 (경고는 plan 이 이미 냈다).
        for name in [
            n for n, sp in list(fields.items())
            if isinstance(sp, dict) and sp.get("from") in ("const", "doc")
            and n not in change["fixed"] and n not in change["doc"]
        ]:
            del fields[name]
            dirty = True

        for name, spec in change["fixed"].items():
            dirty |= _set(fields, name, _flow(_field_entry(spec, "const")))
        for name, spec in change["doc"].items():
            dirty |= _set(fields, name, _flow(_field_entry(spec, "doc")))

        if not dirty:
            continue                      # 바뀐 게 없으면 파일을 건드리지 않는다

        buf = path.with_suffix(".yaml.tmp")
        with buf.open("w", encoding="utf-8") as f:
            y.dump(doc, f)
        buf.replace(path)
        touched.append(str(path.relative_to(ROOT)))

    touched += _apply_brands(sheets.get(S_BRAND, []))
    return touched


def _apply_brands(rows: list[dict]) -> list[str]:
    """`set_keys` 로 코드 묶음씩 저장한다 — 참조표의 행 순서가 판정 우선순위다."""
    grouped: dict[tuple[str, str], list] = {}
    for row in rows:
        kunnr, zbrand, text = row.get("고객코드"), row.get("브랜드코드"), row.get("발주서 원문 문구")
        if not (kunnr and zbrand and text):
            continue
        grouped.setdefault((kunnr, zbrand), []).append(
            brand_store.BrandKey(kunnr=kunnr, zbrand=zbrand,
                                 match=row.get("비교", "contains"),
                                 text=text, note=row.get("비고", ""))
        )

    current = brand_store.load_keys(MASTERS)
    before: dict[tuple[str, str], list] = {}
    for k in current:
        before.setdefault((k.kunnr, k.zbrand), []).append(k)

    for pair in set(before) - set(grouped):     # 엑셀에서 통째로 빠진 코드는 지운다
        grouped[pair] = []

    changed = 0
    for (kunnr, zbrand), keys in grouped.items():
        was = [(k.text, k.match, k.note) for k in before.get((kunnr, zbrand), [])]
        now = [(k.text, k.match, k.note) for k in keys]
        if was == now:
            continue                            # 같으면 파일을 다시 쓰지 않는다
        brand_store.set_keys(MASTERS, kunnr, zbrand, keys)
        changed += 1
    return [f"masters/refs/brand_keys.csv ({changed}개 코드)"] if changed else []


def main() -> int:
    parser = argparse.ArgumentParser(description="거래처 마스터 엑셀 → masters/")
    parser.add_argument("excel", help="master_export.py 로 뽑은 엑셀")
    parser.add_argument("--dry-run", action="store_true", help="계획만 보고 쓰지 않는다")
    parser.add_argument("--masters", default="",
                        help="마스터 폴더 (기본 masters/). 테스트가 사본을 가리킬 때 쓴다")
    parser.add_argument("--yes", action="store_true", help="확인 없이 반영")
    parser.add_argument(
        "--normalize", action="store_true",
        help="값이 그대로여도 전 파일을 다시 쓴다. **한 번만** 쓴다 — 처음 한 번은 "
             "ruamel 서식으로 정리되면서 diff 가 크게 나고, 그 뒤로는 고친 줄만 나온다",
    )
    args = parser.parse_args()

    global MASTERS, CUSTOMERS
    if args.masters:
        MASTERS = Path(args.masters)
        CUSTOMERS = MASTERS / "customers"

    path = Path(args.excel)
    if not path.exists():
        print(f"파일이 없습니다: {path}")
        return 2

    sheets = read_sheets(path)
    errors = validate(sheets)
    if errors:
        print(f"오류 {len(errors)}건 — 아무것도 반영하지 않았습니다.\n")
        for line in errors:
            print("  ⛔", line)
        return 1

    changes, warnings = plan(sheets)
    print(f"거래처 {len(changes)}곳 · 고정값 {len(sheets.get(S_FIXED, []))} · "
          f"문서매핑 {len(sheets.get(S_DOC, []))} · 브랜드 {len(sheets.get(S_BRAND, []))}")
    for line in warnings:
        print("  ⚠", line)

    if args.dry_run:
        print("\n--dry-run 이라 쓰지 않았습니다.")
        return 0
    if warnings and not args.yes:
        answer = input("\n위 내용으로 반영할까요? [y/N] ").strip().lower()
        if answer != "y":
            print("취소했습니다.")
            return 0

    touched = apply(changes, sheets, normalize=args.normalize)
    print("\n반영했습니다:")
    for name in touched:
        print("  ", name)
    print("\n이어서 검증하세요:  python scripts/validate_masters.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
