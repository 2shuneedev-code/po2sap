"""거래처 마스터 엑셀 ↔ masters/ 왕복.

현업이 엑셀을 고치면 규칙이 바뀌고, 규칙이 바뀌면 진짜 오더가 달라진다.
여기서 지키려는 것은 세 가지다.

  · **왕복이 값을 바꾸지 않는다** — 안 고치고 가져오면 파일을 건드리지 않는다
  · **고친 것만 반영된다** — 산출식·읽기 힌트·주석은 그대로 남는다
  · **틀린 값은 아무것도 쓰지 않고 거부된다** — 절반만 반영되면 아무도 모른다
"""

from __future__ import annotations

import importlib.util
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
import yaml

pytest.importorskip("openpyxl")
pytest.importorskip("ruamel.yaml")

from openpyxl import load_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "scripts"


def load_script(name: str):
    """`scripts/` 는 패키지가 아니라 파일이라 직접 읽어 온다."""
    sys.path.insert(0, str(SCRIPTS))
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def workspace(tmp_path, masters_dir):
    dst = tmp_path / "masters"
    shutil.copytree(masters_dir, dst)
    return dst


def run_export(workspace, out: Path) -> None:
    """**항상 사본을 가리킨다.** 실물 masters/ 를 테스트가 덮어쓰면 안 된다."""
    subprocess.run(
        [sys.executable, str(SCRIPTS / "master_export.py"),
         "--out", str(out), "--masters", str(workspace)],
        cwd=ROOT, check=True, capture_output=True,
    )


def run_import(workspace, out: Path, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(SCRIPTS / "master_import.py"), str(out),
         "--masters", str(workspace), *args],
        cwd=ROOT, capture_output=True, text=True,
    )


def customers(path: Path) -> dict[str, dict]:
    return {
        f.stem: yaml.safe_load(f.read_text(encoding="utf-8"))
        for f in sorted((path / "customers").glob("*.yaml"))
        if not f.stem.startswith("_")
    }


# ── 내보내기 ──────────────────────────────────────────────────────────
def test_export_makes_every_sheet(tmp_path, workspace):
    sheets = load_script("master_sheets")
    out = tmp_path / "마스터.xlsx"
    run_export(workspace, out)

    wb = load_workbook(out)
    for name in sheets.COLUMNS:
        assert name in wb.sheetnames, f"{name} 시트가 없다"
    assert sheets.S_GUIDE in wb.sheetnames


def test_export_carries_every_customer(tmp_path, workspace):
    sheets = load_script("master_sheets")
    out = tmp_path / "마스터.xlsx"
    run_export(workspace, out)

    ws = load_workbook(out)[sheets.S_CUSTOMER]
    codes = {r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    assert codes == {c["meta"]["code"] for c in customers(workspace).values()}


def test_expr_and_hints_are_exported_read_only(tmp_path, workspace):
    """표로 못 옮기는 것도 **보여는 준다** — 안 보이면 값의 출처를 못 찾는다."""
    sheets = load_script("master_sheets")
    out = tmp_path / "마스터.xlsx"
    run_export(workspace, out)
    wb = load_workbook(out)

    expr_rows = [r for r in wb[sheets.S_EXPR].iter_rows(min_row=2, values_only=True) if r[0]]
    hint_rows = [r for r in wb[sheets.S_HINTS].iter_rows(min_row=2, values_only=True) if r[0]]
    assert expr_rows, "산출식 시트가 비었다"
    assert hint_rows, "읽기힌트 시트가 비었다"
    assert sheets.S_EXPR in sheets.READ_ONLY
    assert sheets.S_HINTS in sheets.READ_ONLY


# ── 가져오기 ──────────────────────────────────────────────────────────
def test_unchanged_round_trip_writes_nothing(tmp_path, workspace):
    """★ 안 고치고 가져오면 파일을 건드리지 않아야 한다.

    건드리면 아무것도 안 바꾼 가져오기가 커다란 diff 를 내고, 진짜 규칙 변경이
    그 안에 묻힌다. Git 이력이 곧 규칙 변경 대장이다 (CLAUDE.md §6).
    """
    out = tmp_path / "마스터.xlsx"
    run_export(workspace, out)
    before = {
        f: f.read_bytes()
        for f in sorted((workspace / "customers").glob("*.yaml"))
    }

    result = run_import(workspace, out, "--yes")
    assert result.returncode == 0, result.stdout + result.stderr

    for path, content in before.items():
        assert path.read_bytes() == content, f"{path.name} 이 까닭 없이 바뀌었다"


def test_import_rejects_unknown_send_field(tmp_path, workspace):
    """전송 필드는 `_base` 가 정한다. 화면이 컬럼을 새로 만들 수 없다."""
    sheets = load_script("master_sheets")
    out = tmp_path / "마스터.xlsx"
    run_export(workspace, out)

    wb = load_workbook(out)
    ws = wb[sheets.S_FIXED]
    ws.append(["MSC", "NOPE_FIELD", "X", "", ""])
    wb.save(out)

    result = run_import(workspace, out, "--yes")
    assert result.returncode == 1
    assert "NOPE_FIELD" in result.stdout
    assert "전송 필드가 아닙니다" in result.stdout


def test_import_rejects_unknown_customer_reference(tmp_path, workspace):
    sheets = load_script("master_sheets")
    out = tmp_path / "마스터.xlsx"
    run_export(workspace, out)

    wb = load_workbook(out)
    wb[sheets.S_FIXED].append(["NOSUCH", "ZSHCO", "A", "", ""])
    wb.save(out)

    result = run_import(workspace, out, "--yes")
    assert result.returncode == 1
    assert "NOSUCH" in result.stdout


def test_import_rejects_blank_customer_number(tmp_path, workspace):
    """고객코드가 비면 참조표 조회가 전부 빗나간다 — 조용히 빈 값이 나간다."""
    sheets = load_script("master_sheets")
    out = tmp_path / "마스터.xlsx"
    run_export(workspace, out)

    wb = load_workbook(out)
    ws = wb[sheets.S_CUSTOMER]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "MSC":
            row[2].value = None
    wb.save(out)

    result = run_import(workspace, out, "--yes")
    assert result.returncode == 1
    assert "고객코드" in result.stdout


def test_a_failed_import_writes_nothing(tmp_path, workspace):
    """★ 하나라도 걸리면 아무것도 쓰지 않는다. 절반만 반영되면 아무도 모른다."""
    sheets = load_script("master_sheets")
    out = tmp_path / "마스터.xlsx"
    run_export(workspace, out)

    wb = load_workbook(out)
    for row in wb[sheets.S_CUSTOMER].iter_rows(min_row=2):
        if row[0].value == "MSC":
            row[4].value = "바뀐담당"           # 유효한 변경
    wb[sheets.S_FIXED].append(["MSC", "NOPE_FIELD", "X", "", ""])   # 그리고 결함
    wb.save(out)

    before = {f: f.read_bytes() for f in (workspace / "customers").glob("*.yaml")}
    run_import(workspace, out, "--yes")
    for path, content in before.items():
        assert path.read_bytes() == content, f"{path.name} — 거부됐는데 쓰였다"


def test_numbers_do_not_become_floats(tmp_path):
    """엑셀이 `100249` 를 `100249.0` 으로 돌려주면 참조표 조회가 전부 빗나간다."""
    sheets = load_script("master_sheets")
    assert sheets.as_text(100249.0) == "100249"
    assert sheets.as_text(3200) == "3200"
    assert sheets.as_text(None) == ""
    assert sheets.as_text("  A  ") == "A"


def test_file_types_round_trip():
    sheets = load_script("master_sheets")
    assert sheets.text_to_file_types(sheets.file_types_to_text(["htm", "html"])) == ["htm", "html"]
    assert sheets.text_to_file_types("PDF") == ["pdf"]
    assert sheets.text_to_file_types("htm, html") == ["htm", "html"]
